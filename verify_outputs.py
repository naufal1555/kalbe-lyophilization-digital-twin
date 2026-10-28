import docx
import fitz

def audit_docx(path):
    print(f"=== AUDITING DOCX: {path} ===")
    doc = docx.Document(path)
    full_text = []
    for p in doc.paragraphs:
        full_text.append(p.text)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                full_text.append(cell.text)
    combined = "\n".join(full_text)
    print("Total characters:", len(combined))
    print("Total words:", len(combined.split()))
    
    forbidden = [
        ("--", "Double dash"),
        ("—", "Em dash"),
        ("&mdash;", "HTML mdash"),
        ("&bull;", "HTML bull"),
        ("APA", "External APA label"),
        ("FORMAL DECLARATION", "Loud declaration header")
    ]
    for term, label in forbidden:
        count = combined.count(term)
        if count > 0:
            print(f"FAIL: {label} found {count} times!")
        else:
            print(f"PASS: {label} not found.")

    required = [
        ("CONFIDENTIAL  |  Candidate: Ahmad Fathin Naufal (UGM)  |  PSE Internship Assessment", "Confidentiality string"),
        ("Ahmad Fathin Naufal", "Candidate Name"),
        ("Universitas Gadjah Mada (UGM)", "University"),
        ("Process Simulation Engineering (PSE) Intern", "Applied Role"),
        ("Technical Case Study: Process Simulation & Digital Twin", "Assessment Track"),
        ("REFERENCES", "References section"),
        ("PEER-REVIEWED LITERATURE & REGULATORY STANDARDS", "Peer-reviewed subtitle")
    ]
    for term, label in required:
        if term in combined:
            print(f"PASS: {label} is present.")
        else:
            print(f"FAIL: {label} is MISSING!")

def audit_pdf(path):
    print(f"\n=== AUDITING PDF: {path} ===")
    doc = fitz.open(path)
    print("Total pages:", len(doc))
    combined = ""
    for i, p in enumerate(doc):
        text = p.get_text()
        combined += f"\n[Page {i+1}]\n" + text
        
    forbidden = [
        ("--", "Double dash"),
        ("—", "Em dash"),
        ("&mdash;", "HTML mdash"),
        ("APA", "External APA label"),
        ("FORMAL DECLARATION", "Loud declaration header")
    ]
    for term, label in forbidden:
        count = combined.count(term)
        if count > 0:
            print(f"FAIL: {label} found {count} times!")
        else:
            print(f"PASS: {label} not found.")

    required = [
        ("CONFIDENTIAL  |  Candidate: Ahmad Fathin Naufal (UGM)  |  PSE Internship Assessment", "Confidentiality string"),
        ("Ahmad Fathin Naufal", "Candidate Name"),
        ("Universitas Gadjah Mada (UGM)", "University"),
        ("Process Simulation Engineering (PSE) Intern", "Applied Role"),
        ("Technical Case Study: Process Simulation & Digital Twin", "Assessment Track"),
        ("REFERENCES", "References section"),
        ("PEER-REVIEWED LITERATURE & REGULATORY STANDARDS", "Peer-reviewed subtitle")
    ]
    for term, label in required:
        if term in combined:
            print(f"PASS: {label} is present.")
        else:
            print(f"FAIL: {label} is MISSING!")

if __name__ == "__main__":
    audit_docx("report/Technical_Report.docx")
    audit_pdf("report/Technical_Report.pdf")
